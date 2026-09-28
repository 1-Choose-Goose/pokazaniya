import os
import sqlite3
import tempfile
import unittest
import tkinter as tk
from types import SimpleNamespace
from tkinter import ttk
from unittest.mock import Mock, patch
from decimal import Decimal
from pathlib import Path

from main import (
    App, Database, DuplicateEntry, Entry, center_window, export_excel_workbook, number,
    suggested_reading, valid_date, valid_payment_time,
)


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "data.db"
        self.db = Database(self.path)

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def test_separate_services(self):
        self.db.add(Entry(
            "Вода", "22.09.2026", Decimal("450.24"), "23.09.2026", True,
            "111657", "3205834", Decimal("154"), Decimal("71"),
        ))
        self.db.add(Entry("Газ", "01.09.2026", Decimal("58.25")))
        water = self.db.all("Вода")[0]
        self.assertEqual(water.total_water, Decimal("225"))
        self.assertEqual(water.meter_1, "111657")
        self.assertEqual(water.payment_date, "23.09.2026")
        self.assertFalse(self.db.all("Газ")[0].paid)
        self.assertEqual(len(self.db.all()), 2)

    def test_update_and_delete(self):
        entry_id = self.db.add(Entry("ТКО", "01.09.2026", Decimal("2.13")))
        entry = self.db.get(entry_id)
        entry.paid = True
        entry.payment_date = "23.09.2026"
        self.db.update(entry)
        self.assertTrue(self.db.get(entry_id).paid)
        self.db.delete(entry_id)
        self.assertIsNone(self.db.get(entry_id))

    def test_only_one_gas_entry_is_allowed_per_month(self):
        original = Entry("Газ", "22.09.2026", Decimal("256.32"), paid=True,
                         payment_date="22.09.2026 21:16")
        original.id = self.db.add(original)
        with self.assertRaisesRegex(DuplicateEntry, "за этот месяц уже существует"):
            self.db.add(Entry("Газ", "24.09.2026", Decimal("58.25")))
        original.note = "исправлено"
        self.db.update(original)
        self.assertEqual(len(self.db.all("Газ")), 1)

    def test_identical_non_gas_entry_is_rejected(self):
        entry = Entry("ТКО", "24.09.2026", Decimal("100"))
        self.db.add(entry)
        with self.assertRaisesRegex(DuplicateEntry, "уже сохранена"):
            self.db.add(entry)

    def test_legacy_import_is_idempotent(self):
        self.db.close()
        connection = sqlite3.connect(self.path)
        connection.execute(
            """CREATE TABLE readings (
                id INTEGER PRIMARY KEY, reading_date TEXT,
                cold_water TEXT, hot_water TEXT, electricity_t1 TEXT,
                electricity_t2 TEXT, payment_water TEXT,
                payment_electricity TEXT, payment_gas TEXT,
                payment_other TEXT, is_paid INTEGER, notes TEXT
            )"""
        )
        connection.execute(
            """INSERT INTO readings VALUES
            (1, '22.09.2026', '154', '71', '4084', '1485',
             '450.24', '169.60', '58.25', '2.13', 1, 'архив')"""
        )
        connection.commit()
        connection.close()
        self.db = Database(self.path)
        self.assertEqual(len(self.db.all()), 4)
        self.db.close()
        self.db = Database(self.path)
        self.assertEqual(len(self.db.all()), 4)

    def test_input_validation(self):
        self.assertEqual(number("4 084,5", "День"), Decimal("4084.5"))
        self.assertEqual(
            valid_payment_time("23.09.2026 18:45"), "23.09.2026 18:45"
        )
        self.assertEqual(
            valid_payment_time("23.09.2026 12,55"), "23.09.2026 12:55"
        )
        self.assertEqual(
            valid_payment_time("23.09.2026 12.55"), "23.09.2026 12:55"
        )
        with self.assertRaises(ValueError):
            number("NaN", "Сумма")
        with self.assertRaises(ValueError):
            valid_date("31.02.2026")
        with self.assertRaises(ValueError):
            valid_payment_time("23.09.2026 25:00")

    def test_clipboard_shortcuts_in_every_main_input(self):
        self.db.set_setting("gas_meter", "123456")
        app = App(self.db)
        try:
            app.update()
            app.focus_force()
            checked = 0
            for form in app.forms.values():
                for widget in form.input_widgets.values():
                    widget.focus_force()
                    app.update()
                    widget.delete(0, "end")
                    widget.insert(0, "12345")
                    for keycode in (65, 67, 88, 86):
                        widget.event_generate(
                            "<KeyPress>", keycode=keycode, state=4, when="now"
                        )
                    self.assertEqual(widget.get(), "12345")
                    checked += 1
            self.assertGreaterEqual(checked, 20)
        finally:
            app.destroy()

    def test_clipboard_shortcuts_in_settings_and_time_picker(self):
        app = App(self.db)
        try:
            app.update()
            app.edit_meter_settings()
            settings = next(
                widget for widget in app.winfo_children()
                if isinstance(widget, tk.Toplevel)
            )
            settings.update()
            stack = [settings]
            settings_inputs = []
            while stack:
                widget = stack.pop()
                stack.extend(widget.winfo_children())
                if isinstance(widget, ttk.Entry):
                    settings_inputs.append(widget)
            self.assertEqual(len(settings_inputs), 8)
            for widget in settings_inputs:
                widget.focus_force()
                settings.update()
                widget.delete(0, "end")
                widget.insert(0, "12345")
                for code in (65, 67, 88, 86):
                    widget.event_generate(
                        "<KeyPress>", keycode=code, state=4, when="now"
                    )
                self.assertEqual(widget.get(), "12345")
            settings.destroy()
            app.choose_date_time(
                app.forms["Вода"].fields["payment_date"], with_time=True
            )
            picker = next(
                widget for widget in app.winfo_children()
                if isinstance(widget, tk.Toplevel)
            )
            picker.update()
            stack = [picker]
            time_inputs = []
            while stack:
                widget = stack.pop()
                stack.extend(widget.winfo_children())
                if isinstance(widget, tk.Spinbox):
                    time_inputs.append(widget)
            self.assertEqual(len(time_inputs), 2)
            for widget in time_inputs:
                widget.focus_force()
                picker.update()
                widget.delete(0, "end")
                widget.insert(0, "12")
                for code in (65, 67, 88, 86):
                    widget.event_generate(
                        "<KeyPress>", keycode=code, state=4, when="now"
                    )
                self.assertEqual(widget.get(), "12")
            picker.destroy()
        finally:
            app.destroy()

    def test_empty_2025_and_year_arrows(self):
        app = App(self.db)
        try:
            app.update()
            self.assertIn("2025", app.available_years)
            app._select_year("2026")
            app.year_previous.event_generate("<Button-1>", when="now")
            self.assertEqual(app.year_var.get(), "2025")
            self.assertEqual(len(app.tree.get_children()), 12)
            app.year_next.event_generate("<Button-1>", when="now")
            self.assertEqual(app.year_var.get(), "2026")
        finally:
            app.destroy()

    def test_save_shortcut_from_time_field(self):
        app = App(self.db)
        try:
            app.update()
            form = app.forms["Газ"]
            form.fields["amount"].set("58,25")
            time_widget = form.input_widgets["payment_time"]
            time_widget.focus_force()
            app.update()
            time_widget.event_generate(
                "<KeyPress>", keycode=83, state=4, when="now"
            )
            records = self.db.all("Газ")
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].amount, Decimal("58.25"))
        finally:
            app.destroy()

    def test_successful_save_requests_one_cloud_upload(self):
        app = App(self.db)
        app.cloud = SimpleNamespace(saved=Mock())
        try:
            app.forms["Газ"].fields["amount"].set("58,25")
            app.save_all()
            app.cloud.saved.assert_called_once_with()
            self.assertEqual(len(self.db.all("Газ")), 1)
        finally:
            app.destroy()

    def test_invalid_save_does_not_request_cloud_upload(self):
        app = App(self.db)
        app.cloud = SimpleNamespace(saved=Mock())
        try:
            app.forms["Газ"].fields["period"].set("не дата")
            app.forms["Газ"].fields["amount"].set("58,25")
            with patch("main.messagebox.showerror"):
                app.save_all()
            app.cloud.saved.assert_not_called()
            self.assertEqual(self.db.all("Газ"), [])
        finally:
            app.destroy()

    def test_average_consumption_suggestion_is_only_a_hint(self):
        for month, hot in ((1, 31), (2, 51), (3, 71)):
            self.db.add(Entry(
                "Вода", f"20.{month:02d}.2026", None,
                value_1=Decimal("100"), value_2=Decimal(hot),
            ))
        self.assertEqual(
            suggested_reading(
                self.db.all(), "Вода", "value_2", "20.04.2026"
            ),
            Decimal("91.000"),
        )
        app = App(self.db)
        try:
            form = app.forms["Вода"]
            form.fields["period"].set("20.04.2026")
            app.update()
            self.assertEqual(form.suggestion_labels["value_2"].cget("text"), "91")
            self.assertTrue(form.suggestion_labels["value_2"].place_info())
            self.assertEqual(form.fields["value_2"].get(), "")
            form.fields["value_2"].set("90")
            self.assertFalse(form.suggestion_labels["value_2"].place_info())
        finally:
            app.destroy()

    def test_monthly_total_and_unpaid_or_missing_rows(self):
        water_id = self.db.add(Entry(
            "Вода", "22.09.2026", Decimal("450.24"),
            value_1=Decimal("154"), value_2=Decimal("71"),
        ))
        self.db.add(Entry(
            "Газ", "22.09.2026", Decimal("58.25"),
            payment_date="23.09.2026 12:55", paid=True,
        ))
        app = App(self.db)
        try:
            app._select_year("2026")
            month = app.tree.item("month:9")
            self.assertEqual(month["values"][3], "58,25 ₽")
            self.assertEqual(month["values"][-1], "—")
            self.assertIn("attention", app.tree.item(f"entry:{water_id}")["tags"])
            self.assertIn("attention", app.tree.item("missing:9:4")["tags"])
            self.assertEqual(len(app.tree.get_children("month:1")), 0)
            self.assertEqual(
                app.tree.item("missing:9:4")["values"][-1], "Не заполнено"
            )
            for column in ("#0", *app.tree["columns"]):
                self.assertEqual(str(app.tree.column(column, "anchor")), "center")
        finally:
            app.destroy()

    def test_open_month_loads_all_services_without_overwriting_silently(self):
        water_id = self.db.add(Entry(
            "Вода", "22.09.2026", Decimal("450.24"),
            value_1=Decimal("154"), value_2=Decimal("71"),
        ))
        gas_id = self.db.add(Entry(
            "Газ", "21.09.2026", Decimal("58.25"),
        ))
        app = App(self.db)
        try:
            app._select_year("2026")
            app.update()
            app.tree.see("month:9")
            app.update()
            month_box = app.tree.bbox("month:9")
            self.assertTrue(month_box)
            app._open_history_row(SimpleNamespace(y=month_box[1] + month_box[3] // 2))
            self.assertEqual(app.forms["Вода"].edit_id, water_id)
            self.assertEqual(app.forms["Газ"].edit_id, gas_id)
            self.assertEqual(app.forms["Вода"].fields["value_1"].get(), "154")
            self.assertEqual(
                app.forms["Домофон"].fields["period"].get(), "01.09.2026"
            )
            with patch("main.messagebox.askyesno", return_value=False):
                app._load_month(8)
            self.assertEqual(app.forms["Вода"].edit_id, water_id)
            with patch("main.messagebox.askyesno", return_value=True):
                app._load_month(8)
            self.assertIsNone(app.forms["Вода"].edit_id)
            self.assertEqual(
                app.forms["Вода"].fields["period"].get(), "01.08.2026"
            )
        finally:
            app.destroy()

    def test_gas_reading_field_requires_meter_number(self):
        app = App(self.db)
        try:
            gas = app.forms["Газ"]
            self.assertEqual(str(gas.input_widgets["value_1"]["state"]), "disabled")
            self.db.set_setting("gas_meter", "123456")
            gas.refresh_meter_labels()
            self.assertEqual(str(gas.input_widgets["value_1"]["state"]), "normal")
            gas.fields["value_1"].set("42")
            self.db.set_setting("gas_meter", "")
            gas.refresh_meter_labels()
            self.assertEqual(str(gas.input_widgets["value_1"]["state"]), "disabled")
            self.assertEqual(gas.fields["value_1"].get(), "")
        finally:
            app.destroy()

    def test_only_current_month_is_expanded(self):
        from datetime import date

        app = App(self.db)
        try:
            today = date.today()
            app._select_year(str(today.year))
            for month in range(1, 13):
                self.assertEqual(
                    bool(app.tree.item(f"month:{month}", "open")),
                    month == today.month,
                )
            app._select_year("2025")
            self.assertTrue(all(
                not bool(app.tree.item(f"month:{month}", "open"))
                for month in range(1, 13)
            ))
        finally:
            app.destroy()

    def test_clear_all_resets_forms_but_preserves_history(self):
        entry_id = self.db.add(Entry("Газ", "21.09.2026", Decimal("58.25")))
        app = App(self.db)
        try:
            app.forms["Газ"].load(self.db.get(entry_id))
            app.forms["ТКО"].fields["amount"].set("100")
            with patch("main.messagebox.askyesno", return_value=False):
                app.clear_all()
            self.assertEqual(app.forms["Газ"].edit_id, entry_id)
            with patch("main.messagebox.askyesno", return_value=True):
                app.clear_all()
            self.assertTrue(all(form.edit_id is None for form in app.forms.values()))
            self.assertEqual(app.forms["ТКО"].fields["amount"].get(), "")
            self.assertIsNotNone(self.db.get(entry_id))
        finally:
            app.destroy()

    def test_excel_export_contains_full_data_and_formatted_summary(self):
        from openpyxl import load_workbook

        self.db.set_setting("water_cold_meter", "00111657")
        self.db.set_setting("water_cold_date", "08.07.2027")
        self.db.add(Entry(
            "Вода", "22.09.2026", Decimal("450.24"),
            payment_date="23.09.2026 12:55", paid=True,
            meter_1="00111657", meter_2="3205834",
            value_1=Decimal("154"), value_2=Decimal("71"),
            note="Квитанция",
        ))
        self.db.add(Entry("Газ", "22.09.2026", Decimal("58.25")))
        output = Path(self.temp.name) / "export.xlsx"
        export_excel_workbook(output, self.db)
        book = load_workbook(output)
        self.assertEqual(book.sheetnames, ["Сводка", "Все записи", "Счётчики"])
        journal = book["Все записи"]
        self.assertEqual(journal["F5"].value, "00111657")
        self.assertEqual(journal["F5"].data_type, "s")
        self.assertEqual(journal["J5"].value, 225)
        self.assertEqual(journal["K5"].value, 450.24)
        self.assertEqual(journal["M6"].value, "Не оплачено")
        self.assertEqual(book["Счётчики"]["C5"].value, "00111657")
        self.assertEqual(book["Сводка"]["J26"].value, "=SUM(C26:I26)")
        self.assertEqual(journal.freeze_panes, "C5")

    def test_icon_uses_prepared_sizes_instead_of_tk_subsampling(self):
        app = App(self.db)
        try:
            self.assertEqual(app._small_icon.width(), 40)
            self.assertEqual([icon.width() for icon in app._icons], [16, 32, 48, 256])
        finally:
            app.destroy()

    @unittest.skipIf(
        os.environ.get("GITHUB_ACTIONS") == "true",
        "Hosted Windows runner принудительно сдвигает высокие Tk-окна",
    )
    def test_main_and_custom_dialogs_are_centered(self):
        app = App(self.db)
        try:
            app.update_idletasks()
            # Windows themes can include an asymmetric 2–3 px non-client frame
            # in the reported coordinates even when the requested geometry is exact.
            tolerance = 6
            self.assertLessEqual(abs(app.winfo_x() * 2 + app.winfo_width() - app.winfo_screenwidth()), tolerance)
            self.assertLessEqual(abs(app.winfo_y() * 2 + app.winfo_height() - app.winfo_screenheight()), tolerance)
            dialog = tk.Toplevel(app)
            dialog.withdraw()
            ttk.Label(dialog, text="Проверка центра", padding=20).pack()
            center_window(dialog)
            dialog.deiconify()
            dialog.update_idletasks()
            self.assertLessEqual(
                abs(dialog.winfo_x() * 2 + dialog.winfo_width() - dialog.winfo_screenwidth()),
                tolerance,
            )
            self.assertLessEqual(
                abs(dialog.winfo_y() * 2 + dialog.winfo_height() - dialog.winfo_screenheight()),
                tolerance,
            )
            dialog.destroy()
        finally:
            app.destroy()


if __name__ == "__main__":
    unittest.main()
